"""Build the monthly commercial management pack as a formatted Excel workbook.

The dashboard is the self-service surface; this workbook is the artifact that
actually circulates in a commercial team - it is attached to the monthly
business review, it is read on a phone, and it is the thing a regional manager
forwards. It is generated from the warehouse on every run, so it cannot drift
from the dashboard.

Sheets
    Scorecard         headline KPIs vs prior month, last year and plan
    Plan vs Actual    country and region variance over the plan horizon
    Price-Volume-Mix  three-way revenue bridge for three comparisons
    Forecast          backtested outlook to year end, with the model comparison
    Data Quality      validation results and the reconciliation statement
    Definitions       metric definitions and stated assumptions
"""

from __future__ import annotations

import argparse
from pathlib import Path

import duckdb
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

INK = "20242C"
ACCENT = "2457A7"
HEADER_FILL = PatternFill("solid", fgColor="1F3864")
SUBHEAD_FILL = PatternFill("solid", fgColor="D9E2F3")
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
AMBER_FILL = PatternFill("solid", fgColor="FFEB9C")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CURRENCY = '#,##0;[Red](#,##0)'
CURRENCY_2DP = '#,##0.00;[Red](#,##0.00)'
PERCENT = '0.0%;[Red]-0.0%'
NUMBER = '#,##0'


def write_frame(
    worksheet,
    frame: pd.DataFrame,
    start_row: int,
    title: str | None = None,
    formats: dict[str, str] | None = None,
) -> int:
    """Write a dataframe as a formatted block. Returns the next free row."""
    row = start_row
    if title:
        cell = worksheet.cell(row=row, column=1, value=title)
        cell.font = Font(bold=True, size=12, color=INK)
        row += 2

    for index, column in enumerate(frame.columns, start=1):
        cell = worksheet.cell(row=row, column=index, value=str(column))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    header_row = row
    row += 1

    for _, record in frame.iterrows():
        for index, column in enumerate(frame.columns, start=1):
            value = record[column]
            if pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                value = value.item()
            cell = worksheet.cell(row=row, column=index, value=value)
            cell.border = BORDER
            if formats and column in formats:
                cell.number_format = formats[column]
        row += 1

    for index, column in enumerate(frame.columns, start=1):
        longest = max(
            [len(str(column))] + [len(str(v)) for v in frame[column].head(60).tolist()]
        )
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(12, longest + 3), 42
        )

    worksheet.freeze_panes = worksheet.cell(row=header_row + 1, column=1)
    return row + 2


def build_scorecard(writer, connection) -> None:
    frame = connection.execute(
        """
        SELECT kpi_name AS "KPI",
               reporting_month AS "Month",
               actual_value AS "Actual",
               prior_month_value AS "Prior month",
               mom_variance_rate AS "MoM %",
               last_year_value AS "Last year",
               yoy_variance_rate AS "YoY %",
               plan_value AS "Plan",
               plan_achievement_rate AS "vs Plan %",
               plan_rag_status AS "Status"
        FROM mart.commercial_scorecard ORDER BY kpi_order
        """
    ).fetchdf()
    sheet = writer.book.create_sheet("Scorecard")
    reporting_month = pd.to_datetime(frame["Month"].iloc[0]).strftime("%B %Y")
    header = sheet.cell(row=1, column=1, value=f"Commercial Scorecard - {reporting_month}")
    header.font = Font(bold=True, size=14, color=ACCENT)
    sheet.cell(
        row=2,
        column=1,
        value="Latest complete month. Plan targets exist for Net Sales and Units only; "
        "other KPIs are tracked against prior month and last year.",
    ).font = Font(italic=True, size=9, color="595959")

    next_row = write_frame(
        sheet,
        frame,
        start_row=4,
        formats={
            "Actual": NUMBER,
            "Prior month": NUMBER,
            "Last year": NUMBER,
            "Plan": NUMBER,
            "MoM %": PERCENT,
            "YoY %": PERCENT,
            "vs Plan %": PERCENT,
        },
    )

    # write_frame was called without a title, so the header sits on row 4 and
    # the first data row is row 5.
    status_column = get_column_letter(len(frame.columns))
    span = f"{status_column}5:{status_column}{4 + len(frame)}"
    for value, fill in (("Green", GREEN_FILL), ("Amber", AMBER_FILL), ("Red", RED_FILL)):
        sheet.conditional_formatting.add(
            span,
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill),
        )

    trend = connection.execute(
        """
        SELECT month_start AS "Month", net_sales_value AS "Net sales",
               net_profit_value AS "Net profit"
        FROM mart.sales_monthly
        WHERE is_complete_month = 1 AND month_start >= DATE '2022-01-01'
        ORDER BY month_start
        """
    ).fetchdf()
    chart_start = write_frame(sheet, trend, next_row, "Net sales and profit trend",
                              formats={"Net sales": NUMBER, "Net profit": NUMBER})
    chart = LineChart()
    chart.title = "Monthly net sales and net profit"
    chart.height, chart.width = 8, 20
    data_first = next_row + 2
    data = Reference(sheet, min_col=2, max_col=3, min_row=data_first,
                     max_row=data_first + len(trend))
    categories = Reference(sheet, min_col=1, min_row=data_first + 1,
                           max_row=data_first + len(trend))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, f"A{chart_start}")


def build_plan_vs_actual(writer, connection) -> None:
    region = connection.execute(
        """
        SELECT region AS "Region",
               EXTRACT(year FROM plan_year)::INTEGER AS "Plan year",
               actual_net_sales_usd AS "Actual USD",
               plan_net_sales_usd AS "Plan USD",
               variance_usd AS "Variance USD",
               achievement_rate AS "Achievement",
               months_without_plan AS "Months with no plan"
        FROM mart.plan_vs_actual_region ORDER BY "Plan year", "Region"
        """
    ).fetchdf()
    country = connection.execute(
        """
        SELECT canonical_country AS "Country", region AS "Region",
               SUM(actual_net_sales_usd) AS "Actual USD",
               SUM(plan_net_sales_usd) AS "Plan USD",
               SUM(actual_net_sales_usd) - SUM(plan_net_sales_usd) AS "Variance USD",
               SUM(actual_net_sales_usd) / NULLIF(SUM(plan_net_sales_usd), 0) AS "Achievement"
        FROM mart.plan_vs_actual_monthly
        WHERE is_complete_month = 1 AND in_plan_horizon = 1
          AND plan_net_sales_usd IS NOT NULL AND actual_net_sales_usd IS NOT NULL
        GROUP BY 1, 2 ORDER BY "Achievement"
        """
    ).fetchdf()

    sheet = writer.book.create_sheet("Plan vs Actual")
    sheet.cell(row=1, column=1, value="Plan versus Actual").font = Font(
        bold=True, size=14, color=ACCENT
    )
    sheet.cell(
        row=2, column=1,
        value="Complete months inside the plan horizon only. Achievement below 1.00 "
              "means the market landed under its submitted target.",
    ).font = Font(italic=True, size=9, color="595959")

    money = {"Actual USD": NUMBER, "Plan USD": NUMBER, "Variance USD": NUMBER,
             "Achievement": PERCENT}
    next_row = write_frame(sheet, region, 4, "By region and plan year", money)
    country_header = next_row + 1
    write_frame(sheet, country, next_row, "By market, ranked by achievement", money)

    achievement_column = get_column_letter(len(country.columns))
    span = f"{achievement_column}{country_header + 2}:{achievement_column}{country_header + 2 + len(country)}"
    sheet.conditional_formatting.add(
        span, CellIsRule(operator="lessThan", formula=["0.95"], fill=RED_FILL)
    )
    sheet.conditional_formatting.add(
        span, CellIsRule(operator="greaterThanOrEqual", formula=["1.0"], fill=GREEN_FILL)
    )


def build_pvm(writer, connection) -> None:
    totals = connection.execute(
        """
        SELECT comparison_label AS "Comparison",
               SUM(revenue_prior) AS "Prior revenue",
               SUM(revenue_current) AS "Current revenue",
               SUM(total_variance) AS "Total variance",
               SUM(volume_effect) AS "Volume effect",
               SUM(mix_effect) AS "Mix effect",
               SUM(price_effect) AS "Price effect",
               SUM(units_prior) AS "Prior units",
               SUM(units_current) AS "Current units"
        FROM mart.price_volume_mix GROUP BY 1, comparison_order ORDER BY comparison_order
        """
    ).fetchdf()
    detail = connection.execute(
        """
        SELECT comparison_label AS "Comparison", department AS "Department",
               category AS "Category",
               units_prior AS "Prior units", units_current AS "Current units",
               asp_prior AS "Prior ASP", asp_current AS "Current ASP",
               total_variance AS "Total variance", volume_effect AS "Volume",
               mix_effect AS "Mix", price_effect AS "Price"
        FROM mart.price_volume_mix
        WHERE comparison_label LIKE 'YTD%'
        ORDER BY total_variance DESC
        """
    ).fetchdf()

    sheet = writer.book.create_sheet("Price-Volume-Mix")
    sheet.cell(row=1, column=1, value="Price-Volume-Mix bridge").font = Font(
        bold=True, size=14, color=ACCENT
    )
    sheet.cell(
        row=2, column=1,
        value="Volume + Mix + Price sums exactly to Total variance. Volume is growth at "
              "prior mix and prior price; Mix is the shift between categories at prior "
              "price; Price is the ASP movement at current volume.",
    ).font = Font(italic=True, size=9, color="595959")

    money = {c: NUMBER for c in ["Prior revenue", "Current revenue", "Total variance",
                                 "Volume effect", "Mix effect", "Price effect",
                                 "Prior units", "Current units", "Volume", "Mix",
                                 "Price"]}
    money.update({"Prior ASP": CURRENCY_2DP, "Current ASP": CURRENCY_2DP})
    next_row = write_frame(sheet, totals, 4, "Bridge totals", money)

    chart_anchor = next_row
    chart = BarChart()
    chart.type = "col"
    chart.title = "Revenue bridge by effect"
    chart.height, chart.width = 8, 18
    data = Reference(sheet, min_col=5, max_col=7, min_row=6, max_row=6 + len(totals))
    categories = Reference(sheet, min_col=1, min_row=7, max_row=6 + len(totals))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    sheet.add_chart(chart, f"A{chart_anchor}")

    write_frame(sheet, detail, chart_anchor + 18,
                "YTD Jan-Apr 2024 vs 2023, by category", money)


def build_forecast(writer, connection) -> None:
    accuracy = connection.execute(
        """
        SELECT model AS "Model", mape AS "MAPE across horizon",
               mape_first_step AS "MAPE 1 month out",
               mape_final_step AS "MAPE 8 months out",
               worst_month_ape AS "Worst fold", bias AS "Mean bias",
               origins AS "Origins", folds AS "Folds"
        FROM mart.forecast_accuracy ORDER BY mape
        """
    ).fetchdf()
    outlook = connection.execute(
        """
        SELECT month_start AS "Month", basis AS "Basis",
               expected_net_sales_usd AS "Expected net sales",
               plan_net_sales_usd AS "Plan",
               gap_to_plan_usd AS "Gap to plan",
               expected_achievement_rate AS "Expected achievement"
        FROM mart.forecast_vs_plan ORDER BY month_start
        """
    ).fetchdf()
    scenarios = connection.execute(
        """
        SELECT scenario AS "Scenario",
               SUM(forecast_net_sales_usd) AS "Net sales",
               SUM(forecast_net_profit_usd) AS "Net profit",
               MAX(assumed_margin_rate) AS "Assumed margin",
               MAX(volume_multiplier) AS "Volume multiplier",
               MAX(narrative) AS "Assumption"
        FROM mart.forecast_scenarios GROUP BY 1 ORDER BY 2 DESC
        """
    ).fetchdf()

    sheet = writer.book.create_sheet("Forecast")
    sheet.cell(row=1, column=1, value="Forecast to year end").font = Font(
        bold=True, size=14, color=ACCENT
    )
    sheet.cell(
        row=2, column=1,
        value="Model chosen by rolling-origin backtest across all 8 steps of the "
              "forecast horizon - not by in-sample fit and not at one month ahead. "
              "24 origins x 8 steps = 192 folds. Two limitations to read with the "
              "numbers: TheLook is synthetic, so its growth curve is smoother than "
              "real demand and this MAPE is optimistic; and the plan it is compared "
              "against is constructed by src/build_plan_source.py, not supplied. "
              "The full audit is in docs/forecast_review.md.",
    ).font = Font(italic=True, size=9, color="595959")

    next_row = write_frame(
        sheet, accuracy, 4, "Model comparison - rolling-origin backtest across the horizon",
        {"MAPE across horizon": PERCENT, "MAPE 1 month out": PERCENT,
         "MAPE 8 months out": PERCENT, "Worst fold": PERCENT, "Mean bias": PERCENT},
    )
    next_row = write_frame(
        sheet, scenarios, next_row, "Scenarios over the remaining plan horizon",
        {"Net sales": NUMBER, "Net profit": NUMBER, "Assumed margin": PERCENT,
         "Volume multiplier": '0.00'},
    )
    write_frame(
        sheet, outlook, next_row, "Month by month: actual where known, forecast beyond",
        {"Expected net sales": NUMBER, "Plan": NUMBER, "Gap to plan": NUMBER,
         "Expected achievement": PERCENT},
    )


def build_quality(writer, connection) -> None:
    validation = connection.execute(
        """
        SELECT check_name AS "Check", layer AS "Layer", severity AS "Severity",
               failure_count AS "Failures", denominator AS "Rows checked",
               status AS "Status", notes AS "What it protects"
        FROM qa.plan_validation ORDER BY status DESC, check_name
        """
    ).fetchdf()
    statement = connection.execute(
        """
        SELECT stream AS "Stream", stage AS "Stage", entity AS "Entity",
               row_count AS "Rows", value_usd AS "Value USD",
               row_delta AS "Row delta", expectation AS "Expectation"
        FROM qa.reconciliation_statement ORDER BY stream, step_order
        """
    ).fetchdf()
    rejects = connection.execute(
        """
        SELECT submitted_country AS "Submitted country", plan_month AS "Plan month",
               currency_code AS "Currency", reject_reason AS "Reject reason"
        FROM qa.plan_rejects ORDER BY reject_reason, submitted_country
        """
    ).fetchdf()

    sheet = writer.book.create_sheet("Data Quality")
    sheet.cell(row=1, column=1, value="Validation and reconciliation").font = Font(
        bold=True, size=14, color=ACCENT
    )
    sheet.cell(
        row=2, column=1,
        value="PASS means zero failures. WARN means the condition was found, "
              "quarantined and reported - it does not stop publication. FAIL stops "
              "the build and there are none.",
    ).font = Font(italic=True, size=9, color="595959")

    next_row = write_frame(sheet, validation, 4, "Commercial-layer validation",
                           {"Failures": NUMBER, "Rows checked": NUMBER})
    next_row = write_frame(sheet, statement, next_row, "End-to-end reconciliation",
                           {"Rows": NUMBER, "Value USD": NUMBER, "Row delta": NUMBER})
    write_frame(sheet, rejects, next_row,
                "Quarantined plan rows to return to the submitting region")


def build_definitions(writer, connection) -> None:
    rows = [
        ("Net Sales", "Item value excluding Cancelled and Returned items.",
         "core.fact_order_item.net_sales_value"),
        ("Net Profit", "Net sales minus unit cost on the same eligible items.",
         "core.fact_order_item.net_profit_value"),
        ("Units", "Count of order-item rows that are neither cancelled nor returned.",
         "core.fact_order_item"),
        ("ASP", "Net sales divided by units, computed within each category and period.",
         "mart.price_volume_mix"),
        ("Volume effect", "(Q1 - Q0) x prior mix share x prior ASP.",
         "mart.price_volume_mix"),
        ("Mix effect", "Q1 x (current share - prior share) x prior ASP.",
         "mart.price_volume_mix"),
        ("Price effect", "Current units x (current ASP - prior ASP).",
         "mart.price_volume_mix"),
        ("Plan", "Regional submission in local currency, converted at a fixed budget "
                 "FX rate set at the open of the planning cycle.",
         "core.fact_plan"),
        ("Achievement", "Actual net sales divided by plan net sales, USD.",
         "mart.plan_vs_actual_monthly"),
        ("Complete month", "Any month before the maximum source month, which is partial.",
         "core.dim_date.is_complete_month"),
        ("Fiscal year", "Group fiscal year starting 1 October. FY2024 = Oct 2023 to Sep 2024.",
         "core.dim_date.fiscal_year"),
    ]
    assumptions = [
        ("Plan data is constructed", "TheLook ships no budget data. The planning file is "
         "generated by src/build_plan_source.py from a stated, reproducible rule and is "
         "labelled as a constructed scenario everywhere it appears."),
        ("Dataset is synthetic", "Patterns reflect the data generator, not a real market. "
         "The transferable value is the method."),
        ("Forecast accuracy is optimistic", "Synthetic growth is smoother than real demand; "
         "expect materially worse MAPE on real retail data."),
        ("Fiscal year start is an assumption", "Set to October in sql/duckdb/02_core_dimensions.sql. "
         "Changing one number re-bases every fiscal column."),
        ("No promotion or discount field", "Price elasticity and promotion lift cannot be "
         "estimated on this dataset. See docs/commercial_analytics.md."),
    ]
    sheet = writer.book.create_sheet("Definitions")
    sheet.cell(row=1, column=1, value="Metric definitions and assumptions").font = Font(
        bold=True, size=14, color=ACCENT
    )
    next_row = write_frame(
        sheet,
        pd.DataFrame(rows, columns=["Metric", "Definition", "Source"]),
        3, "Definitions",
    )
    write_frame(
        sheet,
        pd.DataFrame(assumptions, columns=["Assumption", "Statement"]),
        next_row, "Stated assumptions and limitations",
    )


def run(project_root: Path) -> dict[str, object]:
    database_path = project_root / "artifacts" / "thelook_analytics.duckdb"
    connection = duckdb.connect(str(database_path), read_only=True)
    destination = project_root / "reports" / "commercial_management_pack.xlsx"
    destination.parent.mkdir(parents=True, exist_ok=True)

    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        # ExcelWriter needs at least one sheet before create_sheet is useful.
        pd.DataFrame({"placeholder": []}).to_excel(writer, sheet_name="_tmp", index=False)
        build_scorecard(writer, connection)
        build_plan_vs_actual(writer, connection)
        build_pvm(writer, connection)
        build_forecast(writer, connection)
        build_quality(writer, connection)
        build_definitions(writer, connection)
        del writer.book["_tmp"]

    connection.close()
    summary = {"path": str(destination), "sheets": 6}
    print(summary)
    return summary


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--project-root", type=Path, default=Path(__file__).resolve().parents[1]
    )
    args = parser.parse_args()
    run(args.project_root.resolve())


if __name__ == "__main__":
    main()
